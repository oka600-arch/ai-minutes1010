from datetime import datetime
from io import BytesIO

from docx import Document

from app.schemas import MeetingInfo
from app.services.docx_builder import build_docx
from app.services.dropbox_service import folder_path_for
from app.services.naming import JST, base_name
from app.services.pdf_builder import build_pdf

INFO = MeetingInfo(meeting_name="営業部 定例会議", meeting_datetime=datetime(2026, 10, 8, 10, tzinfo=JST), meeting_tool="zoom", participants=["山田", "佐藤", "田中"])


def test_docx_contains_all_sections_and_todo_table(sample_minutes):
    doc = Document(BytesIO(build_docx(INFO, sample_minutes)))
    text = "\n".join(p.text for p in doc.paragraphs)
    for heading in ("会議情報", "会議概要", "主な議論", "決定事項", "TODO", "未決事項", "次回確認事項"):
        assert f"■ {heading}" in text
    todo = doc.tables[1]
    assert [c.text for c in todo.rows[0].cells] == ["タスク", "担当者", "期限"]
    assert todo.rows[1].cells[1].text == "佐藤"


def test_pdf_is_generated_with_embedded_japanese_font(sample_minutes):
    pdf = build_pdf(INFO, sample_minutes)
    assert pdf.startswith(b"%PDF")
    assert b"NotoSansJP" in pdf


def test_empty_minutes_still_render(sample_minutes):
    empty = sample_minutes.model_copy(update={"discussion": [], "decisions": [], "todos": [], "unresolved_items": [], "next_actions": [], "summary": ""})
    assert build_docx(INFO, empty)
    assert build_pdf(INFO, empty).startswith(b"%PDF")


def test_names_follow_date_and_meeting_name_rule():
    assert base_name("営業/定例:会議", INFO.meeting_datetime) == "2026-10-08_営業_定例_会議"
    assert folder_path_for("/AI議事録", "営業定例会議", INFO.meeting_datetime) == "/AI議事録/2026/10月/2026-10-08_営業定例会議"
    # UTC で前日になる時刻でも、日本時間の日付を使う
    assert base_name("朝会", datetime.fromisoformat("2026-10-07T23:30:00+00:00")) == "2026-10-08_朝会"
