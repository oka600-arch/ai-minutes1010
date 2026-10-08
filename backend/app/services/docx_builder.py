"""議事録の Word ファイルを生成する。"""

from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.schemas import MeetingInfo, MinutesData
from app.services.naming import format_datetime_ja

FONT_NAME = "游ゴシック"
TOOL_LABELS = {"zoom": "Zoom", "google_meet": "Google Meet", "teams": "Microsoft Teams", "other": "その他"}
DISCLAIMER = "本資料はAIが会議の文字起こしから自動生成したベース資料です。正式な議事録は内容を確認のうえ作成してください。"


class DocxBuildError(Exception):
    pass


def _set_font(run, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = FONT_NAME
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT_NAME)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


def _shade(cell, hex_color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    cell._element.get_or_add_tcPr().append(shd)


def _repeat_header(row) -> None:
    """表がページをまたぐとき、見出し行を次ページにも表示する。"""
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _heading(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    _set_font(p.add_run(f"■ {text}"), size=12, bold=True)


def _para(doc, text: str, *, bullet: bool = False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    if bullet:
        p.paragraph_format.left_indent = Pt(12)
        p.paragraph_format.first_line_indent = Pt(-10)
        text = f"・{text}"
    _set_font(p.add_run(text), size=10.5)


def _bullets_or_none(doc, items: list[str]) -> None:
    if not items:
        _para(doc, "なし")
        return
    for item in items:
        _para(doc, item, bullet=True)


def build_docx(info: MeetingInfo, minutes: MinutesData) -> bytes:
    try:
        doc = Document()
        normal = doc.styles["Normal"]
        normal.font.name = FONT_NAME
        normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT_NAME)

        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_font(title.add_run("議事録（AI自動生成ベース資料）"), size=16, bold=True)

        note = doc.add_paragraph()
        note_run = note.add_run(DISCLAIMER)
        _set_font(note_run, size=9)
        note_run.font.color.rgb = RGBColor(0x55, 0x5F, 0x6D)

        # 1. 会議情報
        _heading(doc, "会議情報")
        info_rows = [
            ("会議名", info.meeting_name),
            ("日時", format_datetime_ja(info.meeting_datetime)),
            ("会議ツール", TOOL_LABELS.get(info.meeting_tool, "その他")),
            ("参加者", "、".join(info.participants) if info.participants else "記載なし"),
        ]
        table = doc.add_table(rows=len(info_rows), cols=2)
        table.style = "Table Grid"
        for i, (label, value) in enumerate(info_rows):
            label_cell, value_cell = table.rows[i].cells
            _shade(label_cell, "EEF2F7")
            _set_font(label_cell.paragraphs[0].add_run(label), size=10, bold=True)
            _set_font(value_cell.paragraphs[0].add_run(value), size=10)
            label_cell.width, value_cell.width = Cm(3.2), Cm(12.8)

        # 2. 会議概要
        _heading(doc, "会議概要")
        _para(doc, minutes.summary or "なし")

        # 3. 主な議論
        _heading(doc, "主な議論")
        if not minutes.discussion:
            _para(doc, "なし")
        for i, item in enumerate(minutes.discussion, start=1):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            _set_font(p.add_run(f"{i}. {item.topic}"), size=10.5, bold=True)
            for line in item.content.split("\n"):
                line = line.strip().lstrip("・").strip()
                if line:
                    _para(doc, line, bullet=True)

        # 4. 決定事項
        _heading(doc, "決定事項")
        _bullets_or_none(doc, minutes.decisions)

        # 5〜7. TODO（担当者・期限）
        _heading(doc, "TODO")
        if not minutes.todos:
            _para(doc, "なし")
        else:
            todo_table = doc.add_table(rows=1, cols=3)
            todo_table.style = "Table Grid"
            todo_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for cell, label in zip(todo_table.rows[0].cells, ("タスク", "担当者", "期限")):
                _shade(cell, "1F3A5F")
                run = cell.paragraphs[0].add_run(label)
                _set_font(run, size=10, bold=True)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                cell.paragraphs[0].paragraph_format.keep_with_next = True
            _repeat_header(todo_table.rows[0])
            for todo in minutes.todos:
                cells = todo_table.add_row().cells
                for cell, value in zip(cells, (todo.task, todo.assignee, todo.deadline)):
                    _set_font(cell.paragraphs[0].add_run(value), size=10)
            for row in todo_table.rows:
                for cell, width in zip(row.cells, (Cm(8.5), Cm(3), Cm(4.5))):
                    cell.width = width

        # 8. 未決事項
        _heading(doc, "未決事項")
        _bullets_or_none(doc, minutes.unresolved_items)

        # 9. 次回確認事項
        _heading(doc, "次回確認事項")
        _bullets_or_none(doc, minutes.next_actions)

        buf = BytesIO()
        doc.save(buf)
        return buf.getvalue()
    except Exception as exc:
        raise DocxBuildError("Wordファイルの生成に失敗しました。") from exc
